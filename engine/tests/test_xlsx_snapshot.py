import json
import sqlite3
from contextlib import closing
from pathlib import Path

import openpyxl
import pytest

from scientific_reading.library_service import LibraryService
from scientific_reading.models import PaperMetadata
from scientific_reading.review_service import ReviewService
from scientific_reading.xlsx_snapshot import REVIEW_COLUMNS, XlsxSnapshotService, XLSX_COLUMNS
from scientific_reading.xlsx_user_fields import BASELINE_SHEET


def _seed(root: Path, count: int = 2) -> None:
    service = LibraryService(root)
    for i in range(count):
        service.ingest(PaperMetadata(title=f"中文文献 {i}", authors=[f"作者{i}"], year=2024 + i, journal="期刊", doi=f"10.1000/{i}", pmid=str(100+i), abstract_en="English", abstract_zh="中文摘要"))
    service.close()


@pytest.mark.parametrize("lock_name", ["~$scientific-reading.xlsx", ".~lock.scientific-reading.xlsx#"])
def test_open_spreadsheet_preserves_file_and_defers_import(tmp_path, lock_name):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    assert service.refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    sheet.cell(2, XLSX_COLUMNS.index("用户笔记") + 1, "已保存但仍在编辑")
    workbook.save(service.target)
    workbook.close()
    before = service.target.read_bytes()
    lock = service.target.with_name(lock_name)
    lock.write_text("office owner", encoding="utf-8")
    for action in (service.import_user_fields, service.refresh):
        result = action()
        assert result["status"] == "pending"
        assert result["error"]["code"] == "xlsx_in_use"
        assert service.target.read_bytes() == before
    with sqlite3.connect(tmp_path / "library.sqlite") as connection:
        assert not connection.execute("SELECT user_notes FROM items").fetchone()[0]
    lock.unlink()
    assert service.refresh()["status"] == "success"
    with sqlite3.connect(tmp_path / "library.sqlite") as connection:
        assert connection.execute("SELECT user_notes FROM items").fetchone()[0] == "已保存但仍在编辑"


def _seed_ready_reader_with_assets(root: Path) -> tuple[str, str]:
    service = LibraryService(root)
    try:
        paper_id = service.ingest(
            PaperMetadata(title="图表索引论文", authors=["作者甲"], year=2026)
        )["paper_id"]
    finally:
        service.close()
    source_sha = "a" * 64
    generation = source_sha[:16]
    paper_root = root / "papers" / paper_id
    generation_root = paper_root / "generations" / generation
    (paper_root / "source.pdf").parent.mkdir(parents=True, exist_ok=True)
    (paper_root / "source.pdf").write_bytes(b"%PDF-1.4\nfixture\n%%EOF\n")
    figure = generation_root / "parsed" / "mineru" / "images" / "figure.jpg"
    table_image = generation_root / "parsed" / "mineru" / "tables" / "table.jpg"
    table_html = generation_root / "parsed" / "mineru" / "tables" / "table.html"
    reader = generation_root / "reading" / "reader.html"
    for path in (figure, table_image, table_html, reader):
        path.parent.mkdir(parents=True, exist_ok=True)
    figure.write_bytes(b"figure")
    table_image.write_bytes(b"table-image")
    table_html.write_text("<table><tr><td>1</td></tr></table>", encoding="utf-8")
    reader.write_text(
        '<html><body><div id="block-figure-caption"></div>'
        '<div id="block-table-caption"></div></body></html>',
        encoding="utf-8",
    )
    (generation_root / "manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "assets": [
                    {
                        "asset_id": "figure-1",
                        "kind": "figure",
                        "page": 1,
                        "relative_path": "parsed/mineru/images/figure.jpg",
                        "caption": "Figure 1. Source caption.",
                    },
                    {
                        "asset_id": "table-1",
                        "kind": "table",
                        "page": 2,
                        "relative_path": "parsed/mineru/tables/table.jpg",
                        "caption": "Table 1. Source caption.",
                    },
                    {
                        "asset_id": "table-1-html",
                        "kind": "table",
                        "page": 2,
                        "relative_path": "parsed/mineru/tables/table.html",
                        "caption": "Table 1. Source caption.",
                    },
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (generation_root / "reading" / "reader-manifest.json").write_text(
        json.dumps(
            {
                "assets": [
                    {"id": "figure-1", "caption_block_id": "figure-caption"},
                    {"id": "table-1", "caption_block_id": "table-caption"},
                    {"id": "table-1-html", "caption_block_id": "table-caption"},
                ]
            }
        ),
        encoding="utf-8",
    )
    translations = generation_root / "reading" / "full" / "translations.json"
    translations.parent.mkdir(parents=True, exist_ok=True)
    translations.write_text(
        json.dumps(
            {
                "translations": [
                    {"block_id": "figure-caption", "translation_zh": "图 1：中文图注。"},
                    {"block_id": "table-caption", "translation_zh": "表 1：中文图注。"},
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    with sqlite3.connect(root / "library.sqlite") as conn:
        conn.execute(
            "INSERT INTO attachments(paper_id,rel_path,sha256,size,validated_at) "
            "VALUES(?, 'source.pdf', ?, 24, '2026-09-05T00:00:00+00:00')",
            (paper_id, source_sha),
        )
        conn.execute(
            "INSERT INTO artifacts(paper_id,kind,rel_path,status,updated_at) "
            "VALUES(?, 'reader', ?, 'ready', '2026-09-05T00:00:00+00:00')",
            (paper_id, f"generations/{generation}/reading/reader.html"),
        )
    return paper_id, generation


def test_snapshot_links_pdf_reader_and_indexes_each_parsed_asset(tmp_path):
    paper_id, generation = _seed_ready_reader_with_assets(tmp_path)

    result = XlsxSnapshotService(tmp_path).refresh()

    assert result["status"] == "success"
    workbook = openpyxl.load_workbook(tmp_path / "library" / "scientific-reading.xlsx")
    assert workbook.sheetnames == ["文献", "图表资产", "整理结论", "_身份", "_个人字段基线", "说明"]
    papers = workbook["文献"]
    paper_headers = {cell.value: cell.column for cell in papers[1]}
    assert papers.cell(2, paper_headers["PDF 路径"]).hyperlink.target == (
        f"../papers/{paper_id}/source.pdf"
    )
    assert papers.cell(2, paper_headers["精读 HTML"]).hyperlink.target == (
        f"../papers/{paper_id}/generations/{generation}/reading/reader.html"
    )
    assert papers.cell(2, paper_headers["图表资产路径"]).value == "查看 2 项"
    assert papers.cell(2, paper_headers["图表资产路径"]).hyperlink.target == (
        "#'图表资产'!A2"
    )
    assets = workbook["图表资产"]
    headers = {cell.value: cell.column for cell in assets[1]}
    assert assets.max_row == 3
    assert assets.cell(2, headers["资产 ID"]).value == "figure-1"
    assert assets.cell(2, headers["PDF 页码"]).value == 1
    assert assets.cell(2, headers["中文图注"]).value == "图 1：中文图注。"
    assert assets.cell(2, headers["图片路径"]).value.endswith("images/figure.jpg")
    assert assets.cell(2, headers["表格 HTML 路径"]).value in (None, "")
    assert assets.cell(3, headers["资产 ID"]).value == "table-1"
    assert assets.cell(3, headers["PDF 页码"]).value == 2
    assert assets.cell(3, headers["图片路径"]).value.endswith("tables/table.jpg")
    assert assets.cell(3, headers["表格 HTML 路径"]).value.endswith("tables/table.html")
    assert assets.cell(3, headers["表格 HTML 路径"]).hyperlink.target.endswith(
        "tables/table.html"
    )
    assert assets.cell(3, headers["精读定位"]).hyperlink.target.endswith(
        "reader.html#block-table-caption"
    )
    workbook.close()


def test_asset_index_uses_active_source_generation_before_reader_is_ready(tmp_path):
    paper_id, _generation = _seed_ready_reader_with_assets(tmp_path)
    old_reader = (
        tmp_path / "papers" / paper_id / "generations" / ("b" * 16)
        / "reading" / "reader.html"
    )
    old_reader.parent.mkdir(parents=True, exist_ok=True)
    old_reader.write_text("<html>旧来源精读</html>", encoding="utf-8")
    with sqlite3.connect(tmp_path / "library.sqlite") as conn:
        conn.execute(
            "UPDATE artifacts SET rel_path=?, status='ready' "
            "WHERE paper_id=? AND kind='reader'",
            (f"generations/{'b' * 16}/reading/reader.html", paper_id),
        )

    assert XlsxSnapshotService(tmp_path).refresh()["status"] == "success"

    workbook = openpyxl.load_workbook(tmp_path / "library" / "scientific-reading.xlsx")
    assets = workbook["图表资产"]
    headers = {cell.value: cell.column for cell in assets[1]}
    assert [assets.cell(row, headers["资产 ID"]).value for row in (2, 3)] == [
        "figure-1", "table-1"
    ]
    assert all(
        not assets.cell(row, headers["精读定位"]).value for row in (2, 3)
    )
    workbook.close()


def test_snapshot_has_fixed_columns_all_rows_and_readme_sheet(tmp_path):
    _seed(tmp_path, 3)
    result = XlsxSnapshotService(tmp_path).refresh()
    assert result["status"] == "success"
    workbook = openpyxl.load_workbook(tmp_path / "library" / "scientific-reading.xlsx")
    assert workbook.sheetnames == ["文献", "图表资产", "整理结论", "_身份", "_个人字段基线", "说明"]
    sheet = workbook["文献"]
    assert tuple(cell.value for cell in next(sheet.iter_rows())) == XLSX_COLUMNS
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 4
    assert rows[1][0] == "中文文献 0"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:{sheet.cell(1, len(XLSX_COLUMNS)).column_letter}{sheet.max_row}"
    assert sheet.column_dimensions["A"].width >= 30
    assert sheet["A1"].fill.fill_type == "solid"
    assert sheet["A2"].alignment.wrap_text is True
    workbook.close()


def test_snapshot_lists_every_confirmed_review_conclusion_with_parent_scope(tmp_path):
    service = LibraryService(tmp_path)
    try:
        paper_id = service.ingest(
            PaperMetadata(title="代谢论文", authors=["作者甲"])
        )["paper_id"]
    finally:
        service.close()
    review = ReviewService(tmp_path)
    try:
        review.bind_session("metabolism-parent", paper_id, "review-child")
        review.confirm_conclusions(
            "review-child",
            [
                {
                    "conclusion_type": "机制",
                    "conclusion_text": "结论一",
                    "evidence_locator": "Figure 2",
                },
                {
                    "conclusion_type": "局限",
                    "conclusion_text": "结论二",
                    "evidence_locator": "Discussion",
                },
            ],
        )
    finally:
        review.close()

    assert XlsxSnapshotService(tmp_path).refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(
        tmp_path / "library" / "scientific-reading.xlsx"
    )
    sheet = workbook["整理结论"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == REVIEW_COLUMNS
    assert len(rows) == 3
    assert rows[1][0:7] == (
        "metabolism-parent", "代谢论文", paper_id, "review-child",
        "机制", "结论一", "Figure 2",
    )
    assert rows[2][4:7] == ("局限", "结论二", "Discussion")
    assert sheet.freeze_panes == "A2"
    assert sheet.protection.sheet is True
    workbook.close()


def test_snapshot_preserves_source_url_from_sqlite(tmp_path):
    source_url = "https://publisher.example/papers/42"
    service = LibraryService(tmp_path)
    service.ingest(
        PaperMetadata(title="URL paper", doi="10.1000/url", source_url=source_url)
    )
    service.close()

    assert XlsxSnapshotService(tmp_path).refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(
        tmp_path / "library" / "scientific-reading.xlsx", read_only=True
    )
    rows = list(workbook["文献"].iter_rows(values_only=True))
    workbook.close()
    headers = {value: index for index, value in enumerate(rows[0])}
    assert rows[1][headers["文献链接"]] == source_url


def test_only_user_columns_are_imported_and_identity_conflicts_are_recorded(tmp_path):
    _seed(tmp_path, 2)
    service = XlsxSnapshotService(tmp_path)
    assert service.refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    first_id = sheet.cell(2, headers["文献 ID"]).value
    original_title = sheet.cell(2, headers["文献名"]).value
    sheet.cell(2, headers["个人思考"], "自己的判断")
    sheet.cell(2, headers["个人理解程度"], "基本理解")
    sheet.cell(2, headers["用户笔记"], "复习图 2")
    sheet.cell(2, headers["文献名"], "不得回写的题名")
    sheet.cell(3, headers["文献 ID"], "changed_identity")
    workbook.save(service.target)
    workbook.close()

    result = service.import_user_fields()
    assert result["updated"] == 0
    assert result["conflicts"] == 1
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    conn.row_factory = sqlite3.Row
    row = conn.execute("SELECT * FROM items WHERE paper_id=?", (first_id,)).fetchone()
    assert row["title"] == original_title
    assert not row["personal_thoughts"]
    assert not row["understanding_level"]
    assert not row["user_notes"]
    conflicts = conn.execute("SELECT value FROM library_meta WHERE key='xlsx_conflicts'").fetchone()[0]
    conn.close()
    assert "identity_changed" in conflicts
    # Repairing the identity allows the full batch; system columns still cannot
    # be imported. Keep the original successful-import protection assertions.
    workbook = openpyxl.load_workbook(service.target)
    workbook["文献"].cell(3, headers["文献 ID"], workbook["_身份"].cell(3, 2).value)
    workbook.save(service.target)
    workbook.close()
    assert service.refresh()["status"] == "success"
    with closing(sqlite3.connect(tmp_path / "library.sqlite")) as conn:
        row = conn.execute(
            "SELECT title, personal_thoughts, understanding_level, user_notes FROM items WHERE paper_id=?",
            (first_id,),
        ).fetchone()
    assert row == (original_title, "自己的判断", "基本理解", "复习图 2")


def test_permission_error_keeps_old_file_and_records_pending_then_retry(tmp_path, monkeypatch):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    target = tmp_path / "library" / "scientific-reading.xlsx"
    old = target.read_bytes()
    real_replace = __import__("os").replace
    def locked(src, dst):
        if str(dst) == str(target):
            raise PermissionError("locked")
        return real_replace(src, dst)
    monkeypatch.setattr("scientific_reading.xlsx_snapshot.os.replace", locked)
    result = service.refresh()
    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_replace_permission_denied"
    assert target.read_bytes() == old
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    meta = dict(conn.execute("SELECT key,value FROM library_meta"))
    conn.close()
    assert meta["xlsx_pending"] == "1"
    assert meta["xlsx_error"] == "xlsx_replace_permission_denied"
    monkeypatch.setattr("scientific_reading.xlsx_snapshot.os.replace", real_replace)
    assert service.refresh()["status"] == "success"
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    meta = dict(conn.execute("SELECT key,value FROM library_meta"))
    conn.close()
    assert meta.get("xlsx_pending") == "0"
    assert meta.get("xlsx_error") in (None, "")


def test_refresh_preserves_conflicting_notes_until_identity_is_repaired(tmp_path):
    _seed(tmp_path, 2)
    service = XlsxSnapshotService(tmp_path)
    assert service.refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    paper_id = sheet.cell(3, headers["文献 ID"]).value
    sheet.cell(2, headers["用户笔记"], "正常行的笔记")
    sheet.cell(3, headers["文献 ID"], "changed_identity")
    sheet.cell(3, headers["用户笔记"], "身份待修复，但笔记必须保留")
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_identity_conflict"
    assert result["updated"] == 0
    assert result["conflicts"] == 1
    assert service.target.read_bytes() == original
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    notes = dict(conn.execute("SELECT paper_id, user_notes FROM items"))
    meta = dict(conn.execute("SELECT key,value FROM library_meta"))
    conn.close()
    assert all(not value for value in notes.values())
    assert not notes[paper_id]
    assert meta["xlsx_pending"] == "1"
    assert meta["xlsx_error"] == "xlsx_identity_conflict"

    workbook = openpyxl.load_workbook(service.target)
    workbook["文献"].cell(3, headers["文献 ID"], paper_id)
    workbook.save(service.target)
    workbook.close()
    assert service.refresh()["status"] == "success"
    workbook = openpyxl.load_workbook(service.target)
    assert workbook["文献"].cell(3, headers["用户笔记"]).value == "身份待修复，但笔记必须保留"
    workbook.close()
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    assert conn.execute("SELECT user_notes FROM items WHERE paper_id=?", (paper_id,)).fetchone()[0] == "身份待修复，但笔记必须保留"
    meta = dict(conn.execute("SELECT key,value FROM library_meta"))
    conn.close()
    assert meta["xlsx_pending"] == "0"
    assert "xlsx_error" not in meta
    assert meta["xlsx_conflicts"] == "[]"


@pytest.mark.parametrize("sheet_name", ["文献", "_身份"])
def test_refresh_preserves_workbook_when_required_sheet_is_missing(tmp_path, sheet_name):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    workbook.create_sheet("个人记录").append(["不能覆盖的个人记录"])
    del workbook[sheet_name]
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_required_sheets_missing"
    assert service.target.read_bytes() == original


def test_refresh_reports_missing_user_columns_without_overwriting_workbook(tmp_path):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    column = XLSX_COLUMNS.index("用户笔记") + 1
    sheet.cell(1, column, "改过名字的笔记列")
    sheet.cell(2, column, "表头错误也不能丢失笔记")
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_user_columns_missing"
    assert service.target.read_bytes() == original


@pytest.mark.parametrize("column", [1, 2])
def test_refresh_preserves_workbook_when_identity_header_is_invalid(tmp_path, column):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    workbook["_身份"].cell(1, column, "invalid_header")
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_identity_columns_invalid"
    assert service.target.read_bytes() == original


def test_refresh_preserves_workbook_when_last_paper_row_is_missing(tmp_path):
    _seed(tmp_path, 2)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    workbook["文献"].delete_rows(3)
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_identity_conflict"
    assert result["conflicts"] == 1
    assert service.target.read_bytes() == original


def test_refresh_preserves_handwritten_note_when_user_header_is_duplicated(tmp_path):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    note_column = XLSX_COLUMNS.index("用户笔记") + 1
    sheet.cell(2, note_column, "不能从重复表头的错误列读取")
    sheet.cell(1, sheet.max_column + 1, "用户笔记")
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_user_columns_ambiguous"
    assert service.target.read_bytes() == original


def test_refresh_does_not_misattribute_notes_when_identity_row_is_duplicated(tmp_path):
    _seed(tmp_path, 2)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    first_id = sheet.cell(2, headers["文献 ID"]).value
    second_id = sheet.cell(3, headers["文献 ID"]).value
    sheet.cell(2, headers["文献 ID"], second_id)
    sheet.cell(2, headers["用户笔记"], "属于第一行的笔记")
    sheet.cell(3, headers["用户笔记"], "属于第二行的笔记")
    workbook["_身份"].append((2, second_id))
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_identity_conflict"
    assert service.target.read_bytes() == original
    conn = sqlite3.connect(tmp_path / "library.sqlite")
    notes = dict(conn.execute("SELECT paper_id, user_notes FROM items"))
    conn.close()
    assert notes[first_id] is None
    assert notes[second_id] is None


def test_refresh_preserves_unreadable_workbook_and_allows_file_repair(tmp_path):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    valid = service.target.read_bytes()
    service.target.write_bytes(b"not an xlsx archive")
    original = service.target.read_bytes()

    result = service.refresh()

    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_read_failed"
    assert service.target.read_bytes() == original
    moved = service.target.with_suffix(".broken")
    service.target.replace(moved)
    moved.replace(service.target)
    service.target.write_bytes(valid)
    assert service.refresh()["status"] == "success"


FIELD_NAMES = ("个人思考", "个人理解程度", "用户笔记")
SQL_FIELDS = "personal_thoughts, understanding_level, user_notes"


def _write_personal(root, values):
    with closing(sqlite3.connect(root / "library.sqlite")) as conn, conn:
        conn.execute("UPDATE items SET personal_thoughts=?, understanding_level=?, user_notes=?", values)


def _read_personal(root):
    with closing(sqlite3.connect(root / "library.sqlite")) as conn:
        return conn.execute(f"SELECT {SQL_FIELDS} FROM items ORDER BY paper_id").fetchall()


def _edit_personal(service, values, row=2):
    workbook = openpyxl.load_workbook(service.target)
    for name, value in zip(FIELD_NAMES, values):
        workbook["文献"].cell(row, XLSX_COLUMNS.index(name) + 1).value = value
    workbook.save(service.target)
    workbook.close()


@pytest.mark.parametrize("excel,database,expected", [
    (("B",)*3, ("D",)*3, ("D",)*3),  # F1: untouched old workbook
    (("E",)*3, ("B",)*3, ("E",)*3),
    (("same",)*3, ("same",)*3, ("same",)*3),
    (("E", "B", "B"), ("B", "D", "B"), ("E", "D", "B")),
    ((None, "", None), ("B",)*3, ("",)*3),
    (("B",)*3, ("",)*3, ("",)*3),
    (("B",)*3, ("B",)*3, ("B",)*3),
    (("=literal", "中文\n换行", "#N/A"), ("B",)*3, ("=literal", "中文\n换行", "#N/A")),
])
def test_three_way_user_fields(tmp_path, excel, database, expected):
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    assert service.refresh()["status"] == "success"
    _edit_personal(service, excel)
    _write_personal(tmp_path, database)
    assert service.refresh()["status"] == "success"
    assert _read_personal(tmp_path) == [expected]
    # A new service (restart) reads the newly published baseline, idempotently.
    assert XlsxSnapshotService(tmp_path).refresh()["status"] == "success"
    assert _read_personal(tmp_path) == [expected]
    workbook = openpyxl.load_workbook(service.target)
    assert workbook[BASELINE_SHEET].sheet_state == "veryHidden"
    workbook.close()


@pytest.mark.parametrize("field", range(3))
def test_content_conflict_preserves_whole_batch_and_retries(tmp_path, field):
    _seed(tmp_path, 2)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    excel = ["B"]*3
    excel[field] = "Excel edit"
    _edit_personal(service, excel)
    _edit_personal(service, ("other row edit", "B", "B"), row=3)
    database = ["B"]*3
    database[field] = "DB edit"
    _write_personal(tmp_path, database)
    before = service.target.read_bytes()
    for _ in range(2):
        result = XlsxSnapshotService(tmp_path).refresh()
        assert result["status"] == "pending"
        assert result["updated"] == 0
        assert any(c.get("field") == SQL_FIELDS.split(", ")[field] for c in result["conflict_details"])
        assert service.target.read_bytes() == before
        assert _read_personal(tmp_path) == [tuple(database)]*2


@pytest.mark.parametrize("different", [False, True])
def test_legacy_workbook_requires_equality(tmp_path, different):
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    workbook = openpyxl.load_workbook(service.target)
    del workbook[BASELINE_SHEET]
    workbook["_身份"]["C1"] = None
    workbook["_身份"]["D1"] = None
    workbook.save(service.target)
    workbook.close()
    if different:
        _edit_personal(service, ("E", "B", "B"))
    original = service.target.read_bytes()
    result = service.refresh()
    assert result["status"] == ("pending" if different else "success")
    assert _read_personal(tmp_path) == [("B",)*3]
    if different:
        assert service.target.read_bytes() == original
        assert result["conflict_details"][0]["code"] == "baseline_missing_difference"
    else:
        workbook = openpyxl.load_workbook(service.target)
        assert BASELINE_SHEET in workbook.sheetnames
        workbook.close()


@pytest.mark.parametrize("damage", ["sheet", "marker", "checksum", "value", "duplicate", "missing", "id", "header", "formula"])
def test_damaged_baseline_never_falls_back_to_import(tmp_path, damage):
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    _edit_personal(service, ("E",)*3)
    workbook = openpyxl.load_workbook(service.target)
    baseline = workbook[BASELINE_SHEET]
    if damage == "sheet":
        del workbook[BASELINE_SHEET]
    elif damage in ("marker", "checksum"):
        workbook["_身份"]["C1" if damage == "marker" else "D1"] = "broken"
    elif damage == "duplicate":
        baseline.append(tuple(cell.value for cell in baseline[2]))
    elif damage == "missing":
        baseline.delete_rows(2)
    else:
        baseline[{"value": "B2", "id": "A2", "header": "A1", "formula": "B2"}[damage]] = "=1" if damage == "formula" else "broken"
    workbook.save(service.target)
    workbook.close()
    original = service.target.read_bytes()
    result = service.refresh()
    assert result["status"] == "pending"
    assert result["updated"] == 0
    assert result["conflict_details"][0]["code"] == "baseline_invalid"
    assert service.target.read_bytes() == original
    assert _read_personal(tmp_path) == [("B",)*3]


def test_import_commit_replace_failure_and_retry_is_idempotent(tmp_path, monkeypatch):
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    _edit_personal(service, ("E",)*3)
    original = service.target.read_bytes()
    real_replace = __import__("os").replace
    def fail(source, target):
        if str(target) == str(service.target):
            raise PermissionError("fixture")
        return real_replace(source, target)
    with monkeypatch.context() as patch:
        patch.setattr("scientific_reading.xlsx_snapshot.os.replace", fail)
        assert service.refresh()["status"] == "pending"
    assert _read_personal(tmp_path) == [("E",)*3]
    assert service.target.read_bytes() == original
    assert XlsxSnapshotService(tmp_path).refresh()["status"] == "success"
    assert _read_personal(tmp_path) == [("E",)*3]


def test_import_rereads_database_after_excel_load(tmp_path, monkeypatch):
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    _edit_personal(service, ("E", "B", "B"))
    original = service.target.read_bytes()
    real_load = openpyxl.load_workbook
    def concurrent_load(*args, **kwargs):
        workbook = real_load(*args, **kwargs)
        _write_personal(tmp_path, ("D",)*3)
        return workbook
    monkeypatch.setattr("scientific_reading.xlsx_snapshot.load_workbook", concurrent_load)
    assert service.refresh()["status"] == "pending"
    assert _read_personal(tmp_path) == [("D",)*3]
    assert service.target.read_bytes() == original


def test_writer_reservation_covers_read_through_commit(tmp_path, monkeypatch):
    import threading
    _seed(tmp_path, 1)
    _write_personal(tmp_path, ("B",)*3)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    _edit_personal(service, ("E",)*3)
    real_connect = sqlite3.connect
    attempts = []
    def attempt_write():
        with closing(real_connect(tmp_path / "library.sqlite", timeout=0)) as conn:
            try:
                conn.execute("UPDATE items SET user_notes='concurrent'")
                conn.commit()
                attempts.append("wrote")
            except sqlite3.OperationalError as error:
                attempts.append(str(error))
    class Connection(sqlite3.Connection):
        def execute(self, sql, *args, **kwargs):
            result = super().execute(sql, *args, **kwargs)
            if sql.startswith("SELECT personal_thoughts,"):
                thread = threading.Thread(target=attempt_write)
                thread.start()
                thread.join(timeout=3)
                assert not thread.is_alive()
            return result
    def connect(*args, **kwargs):
        kwargs["factory"] = Connection
        return real_connect(*args, **kwargs)
    monkeypatch.setattr("scientific_reading.xlsx_snapshot.sqlite3.connect", connect)
    assert service.refresh()["status"] == "success"
    assert attempts == ["database is locked"]
    # A later committed writer is retained on the next refresh.
    _write_personal(tmp_path, ("D",)*3)
    assert service.refresh()["status"] == "success"
    assert _read_personal(tmp_path) == [("D",)*3]


@pytest.mark.parametrize("damage", ["sort", "invalid_identity", "unknown_db"])
def test_identity_errors_block_other_valid_edits(tmp_path, damage):
    _seed(tmp_path, 2)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    _edit_personal(service, ("edit", "edit", "edit"))
    workbook = openpyxl.load_workbook(service.target)
    sheet = workbook["文献"]
    if damage == "sort":
        first = [cell.value for cell in sheet[2]]
        second = [cell.value for cell in sheet[3]]
        for column, (one, two) in enumerate(zip(first, second), 1):
            sheet.cell(2, column).value = two
            sheet.cell(3, column).value = one
    elif damage == "invalid_identity":
        workbook["_身份"].append(("not a row", "unknown"))
    else:
        paper_id = sheet.cell(3, XLSX_COLUMNS.index("文献 ID") + 1).value
        with closing(sqlite3.connect(tmp_path / "library.sqlite")) as conn, conn:
            conn.execute("DELETE FROM items WHERE paper_id=?", (paper_id,))
    workbook.save(service.target)
    workbook.close()
    before = service.target.read_bytes()
    database = _read_personal(tmp_path)
    result = service.refresh()
    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_identity_conflict"
    assert result["updated"] == 0
    assert service.target.read_bytes() == before
    assert _read_personal(tmp_path) == database


def test_overlong_database_note_is_not_silently_truncated(tmp_path):
    _seed(tmp_path, 1)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    before = service.target.read_bytes()
    _write_personal(tmp_path, ("B", "B", "x"*32768))
    result = service.refresh()
    assert result["status"] == "failed"
    assert result["error"]["code"] == "xlsx_snapshot_failed"
    assert service.target.read_bytes() == before
    assert _read_personal(tmp_path) == [("B", "B", "x"*32768)]
